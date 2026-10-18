"""The CBSL collector: index pages, the three spreadsheets, and the collector end to end (no network).

Fixtures (tests/fixtures/cbsl) are the real files downloaded on 2026-10-18, cut down to a few rows with their
layout kept: table 4.04 for 2017 to January 2019, table 4.03 for 2017 to 2020, a few weeks of the auction file, and
only the file links of the two statistical-tables pages.
"""
from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import openpyxl
import pytest

from data_platform.collectors.cbsl import collector as cbsl
from data_platform.collectors.cbsl import index, parse
from data_platform.core.collector import Context
from data_platform.core.http import Response

FIX = Path(__file__).parent / "fixtures" / "cbsl"
D = dt.date


def fixture(name: str) -> bytes:
    return (FIX / name).read_bytes()


def values(parsed, series: str) -> dict:
    return {r["month"]: r["value"] for r in parsed.rows if r["series"] == series}


# --- index pages ------------------------------------------------------------------------------------------------

def test_index_pages_give_the_current_file_and_its_date():
    links = index.parse_index(fixture("monetary-sector.html"), "monetary-sector")
    assert {(x.dataset, x.file_date) for x in links} == {("cbsl_4.04", D(2026, 10, 12)), ("cbsl_4.03", D(2026, 3, 5))}
    assert links[0].url.startswith("https://www.cbsl.gov.lk/sites/default/files/")
    (bills,) = index.parse_index(fixture("fiscal-sector.html"), "fiscal-sector")
    assert bills.dataset == "cbsl_tbill_auctions" and bills.file_date == D(2025, 12, 24)


def test_an_index_page_without_the_file_or_with_two_fails():
    html = fixture("monetary-sector.html").decode()
    with pytest.raises(index.IndexChanged, match="cbsl_4.04"):
        index.parse_index(html.replace("table4.04_", "table4.99_"), "monetary-sector")
    doubled = html.replace("</ul>", '<li><a href="/x/table4.04_20991231.xlsx">Interest Rates - Monthly</a></li></ul>')
    with pytest.raises(index.IndexChanged, match="found 2"):
        index.parse_index(doubled, "monetary-sector")


def test_a_relative_link_is_made_absolute():
    html = '<a href="/sites/default/files/x/table4.04_20270110.xlsx">Interest Rates - Monthly (2003 to Latest)</a>' \
           '<a href="/sites/default/files/x/table4.03_20270110.xlsx">Interest Rates - Annual</a>'
    links = index.parse_index(html, "monetary-sector")
    assert links[0].url == "https://www.cbsl.gov.lk/sites/default/files/x/table4.04_20270110.xlsx"


# --- table 4.04 -------------------------------------------------------------------------------------------------

def test_monthly_table_values_as_published():
    p = parse.parse_monthly_rates(fixture("table4.04_cut.xlsx"))
    assert p.dataset == "cbsl_4.04"
    months = sorted({r["month"] for r in p.rows})
    assert months[0] == D(2017, 1, 1) and months[-1] == D(2019, 1, 1) and len(months) == 25
    # spot checks against the published table (December 2018 row)
    assert values(p, "AWNFDR")[D(2018, 12, 1)] == pytest.approx(11.2682990500378)
    assert values(p, "AWFDR")[D(2018, 12, 1)] == 10.85
    assert values(p, "TB_12M")[D(2018, 12, 1)] == 11.2
    assert values(p, "SDFR")[D(2018, 12, 1)] == 8 and values(p, "SLFR")[D(2018, 12, 1)] == 9
    assert values(p, "FD_12M_MIN")[D(2018, 12, 1)] == 4.53 and values(p, "FD_12M_MAX")[D(2018, 12, 1)] == 15
    # "-" is "no figure": the 3-month bill had no accepted bids at the end of January and December 2018
    assert D(2018, 1, 1) not in values(p, "TB_3M") and D(2018, 12, 1) not in values(p, "TB_3M")
    assert values(p, "TB_3M")[D(2018, 2, 1)] == 8.24
    assert values(p, "OPR") == {}                      # the OPR column exists, with figures from November 2024 only


def _rewrite(body: bytes, edit) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(body))
    edit(wb.active)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _find(ws, text: str):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().startswith(text):
                return c
    raise KeyError(text)


def test_columns_are_found_by_header_not_position():
    def insert_column(ws):
        ws.insert_cols(4)                               # a new column before the policy rates shifts everything
    p = parse.parse_monthly_rates(_rewrite(fixture("table4.04_cut.xlsx"), insert_column))
    assert values(p, "AWFDR")[D(2018, 12, 1)] == 10.85 and values(p, "TB_12M")[D(2018, 12, 1)] == 11.2


def test_a_missing_or_renamed_column_is_an_error():
    def rename(ws):
        _find(ws, "AWNFDR").value = "Something else"
    with pytest.raises(parse.ParseError, match="AWNFDR"):
        parse.parse_monthly_rates(_rewrite(fixture("table4.04_cut.xlsx"), rename))


def test_a_gap_in_the_months_is_an_error():
    def drop_june(ws):
        _find(ws, "June").value = "July"
    with pytest.raises(parse.ParseError, match="without a gap"):
        parse.parse_monthly_rates(_rewrite(fixture("table4.04_cut.xlsx"), drop_june))


def test_text_in_a_number_cell_is_an_error_not_a_blank():
    def garble(ws):
        c = _find(ws, "AWFDR")
        ws.cell(c.row + 5, c.column).value = "11.5 (provisional)"
    with pytest.raises(parse.ParseError, match="not a number"):
        parse.parse_monthly_rates(_rewrite(fixture("table4.04_cut.xlsx"), garble))


def test_a_spreadsheet_that_is_not_one_is_refused():
    with pytest.raises(parse.ParseError, match="not an .xlsx"):
        parse.parse_monthly_rates(b"<html>Service unavailable</html>")
    with pytest.raises(parse.ParseError, match="not table 4.04"):
        parse.parse_monthly_rates(fixture("table4.03_cut.xlsx"))


@pytest.mark.parametrize("cell, expected", [
    ("7.00-11.00", (7.0, 11.0)), ("2.50.-11.75", (2.5, 11.75)), (" 4.53 - 15.00 ", (4.53, 15.0)), (6.4, (6.4, 6.4)),
    ("-", None), ("n.a", None), (None, None)])
def test_ranges_and_missing_marks(cell, expected):
    assert parse.value_range(cell, where="t") == expected


def test_a_backwards_range_is_an_error():
    with pytest.raises(parse.ParseError, match="backwards"):
        parse.value_range("11.00-7.00", where="t")


# --- table 4.03 -------------------------------------------------------------------------------------------------

def test_annual_table():
    p = parse.parse_annual_rates(fixture("table4.03_cut.xlsx"))
    by = {(r["series"], r["year"]): r["value"] for r in p.rows}
    assert {r["year"] for r in p.rows} == {2017, 2018, 2019, 2020}
    assert by[("AWNFDR", 2018)] == pytest.approx(11.2682990500378)      # the same figure as December in 4.04
    assert by[("TB_12M", 2020)] == 5.05 and by[("FD_12M_MIN", 2019)] == 3.55 and by[("FD_12M_MAX", 2019)] == 15


# --- the weekly auction file ------------------------------------------------------------------------------------

def test_auction_file():
    p = parse.parse_tbill_auctions(fixture("tbill_auctions_cut.xlsx"))
    rows = {(r["issue_date"], r["tenor_days"]): r for r in p.rows}
    r = rows[(D(2015, 7, 24), 364)]
    assert r["yield_pct"] == pytest.approx(7.1) and r["yield_without_tax_pct"] == 6.39 and r["accepted_mn"] == 3341
    assert rows[(D(2015, 8, 7), 91)]["yield_pct"] is None            # printed 0: nothing accepted
    assert rows[(D(2015, 8, 7), 91)]["offered_mn"] == 0
    assert not any(d == D(2018, 12, 28) for d, _ in rows)             # "Auction was not held"
    late = rows[(D(2006, 9, 29), 364)]                                 # accepted amounts and yields all "-"
    assert late["accepted_mn"] is None and late["yield_pct"] is None and late["offered_mn"] is not None
    assert "1 auctions not held" in p.note


def test_auction_dates_must_increase():
    def swap(ws):
        dates = [c for row in ws.iter_rows() for c in row if isinstance(c.value, dt.datetime)]
        dates[1].value, dates[2].value = dates[2].value, dates[1].value
    with pytest.raises(parse.ParseError, match="date order"):
        parse.parse_tbill_auctions(_rewrite(fixture("tbill_auctions_cut.xlsx"), swap))


# --- the collector against a fake CBSL website ------------------------------------------------------------------

class FakeCbsl:
    def __init__(self, broken: str | None = None):
        self.broken = broken
        self.requests = []

    def request(self, method, url, params=None):
        self.requests.append(url)
        if url.endswith("/monetary-sector"):
            body = fixture("monetary-sector.html")
        elif url.endswith("/fiscal-sector"):
            body = fixture("fiscal-sector.html")
        elif "table4.04_" in url:
            body = b"<html>error</html>" if self.broken == "4.04" else fixture("table4.04_cut.xlsx")
        elif "table4.03_" in url:
            body = fixture("table4.03_cut.xlsx")
        elif "Treasury_Bill_Auctions_" in url:
            body = fixture("tbill_auctions_cut.xlsx")
        else:
            raise AssertionError(f"unexpected request {url}")
        return Response(method, url, params or {}, 200, body, None, dt.datetime.now(dt.timezone.utc))


def run(conn, store, http):
    return cbsl.CbslCollector(Context(conn, store, http, None)).collect()


@pytest.mark.db
def test_first_run_stores_every_file_and_the_second_fetches_only_the_pages(conn, store):
    http = FakeCbsl()
    result = run(conn, store, http)
    conn.commit()
    assert result.new_data and result.rows_added > 0
    n = {t: conn.execute(f"select count(*) n from rates.{t}").fetchone()["n"]
         for t in ("cbsl_monthly", "cbsl_annual", "tbill_auctions", "source_files")}
    assert n["source_files"] == 3 and all(v > 0 for v in n.values())
    assert conn.execute("select count(*) n from rates.cbsl_monthly_current").fetchone()["n"] == n["cbsl_monthly"]
    again = FakeCbsl()
    second = run(conn, store, again)
    assert not second.new_data and second.rows_added == 0
    assert len(again.requests) == 2                     # only the two index pages: each file is read once


@pytest.mark.db
def test_a_new_version_of_a_table_becomes_current_and_the_old_rows_stay(conn, store):
    run(conn, store, FakeCbsl())
    conn.commit()

    def revised(ws):                                     # the CBSL revises one figure in its next update
        ws.cell(_find(ws, "January").row, _find(ws, "AWFDR").column).value = 99.0     # January 2017
    newer = _rewrite(fixture("table4.04_cut.xlsx"), revised)
    http = FakeCbsl()
    original = http.request

    def with_new_file(method, url, params=None):
        if url.endswith("/monetary-sector"):
            body = fixture("monetary-sector.html").replace(b"table4.04_20261012", b"table4.04_20261110")
            return Response(method, url, {}, 200, body, None, dt.datetime.now(dt.timezone.utc))
        if "table4.04_20261110" in url:
            return Response(method, url, {}, 200, newer, None, dt.datetime.now(dt.timezone.utc))
        return original(method, url, params)

    http.request = with_new_file
    result = run(conn, store, http)
    conn.commit()
    assert result.new_data
    current = conn.execute("select value from rates.cbsl_monthly_current where series = 'AWFDR' and month = '2017-01-01'")
    assert current.fetchone()["value"] == 99.0
    kept = conn.execute("select count(*) n from rates.cbsl_monthly where series = 'AWFDR' and month = '2017-01-01'")
    assert kept.fetchone()["n"] == 2                     # nothing overwritten: both versions are stored
    with pytest.raises(Exception, match="append-only|not allowed|forbid"):
        conn.execute("update rates.cbsl_monthly set value = 1")


@pytest.mark.db
def test_a_file_that_cannot_be_read_is_recorded_and_fails_the_run(conn, store):
    with pytest.raises(RuntimeError, match="Parse failed"):
        run(conn, store, FakeCbsl(broken="4.04"))
    conn.rollback()
    row = conn.execute("select parse_status from rates.source_files where dataset = 'cbsl_4.04'").fetchone()
    assert row["parse_status"] == "failed"
    assert conn.execute("select count(*) n from rates.cbsl_monthly_current").fetchone()["n"] == 0
    assert conn.execute("select count(*) n from rates.cbsl_annual_current").fetchone()["n"] > 0   # the others came in
